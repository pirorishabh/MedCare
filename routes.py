"""DoctorTalk page and API routes."""
import logging, os, json, uuid
from functools import wraps
from flask import Blueprint, request, jsonify, render_template, session, redirect, url_for, send_file
from services import watsonx
from database import (create_user, authenticate_user, record_login, get_db, add_chat_message,
                      get_chat_messages, list_doctors, create_appointment)
from utils import sanitize_input, save_upload, extract_text_from_file, validate_json_body, success_response, error_response
from config import config

logger=logging.getLogger(__name__)
main_bp=Blueprint('main',__name__); api_bp=Blueprint('api',__name__,url_prefix='/api')

def login_required(view):
    @wraps(view)
    def wrapped(*args,**kwargs):
        if not session.get('user_id'):
            if request.path.startswith('/api/'): return error_response('Login required.',401)
            return redirect(url_for('main.login',next=request.path))
        return view(*args,**kwargs)
    return wrapped

def doctor_required(view):
    @wraps(view)
    def wrapped(*args,**kwargs):
        if not session.get('user_id') or session.get('role')!='doctor':
            if request.path.startswith('/api/'): return error_response('Doctor login required.',403)
            return redirect(url_for('main.doctor_login'))
        return view(*args,**kwargs)
    return wrapped

def patient_required(view):
    @wraps(view)
    def wrapped(*args,**kwargs):
        if not session.get('user_id') or session.get('role')!='patient':
            if request.path.startswith('/api/'): return error_response('Patient login required.',403)
            return redirect(url_for('main.login'))
        return view(*args,**kwargs)
    return wrapped

@main_bp.route('/')
def index(): return redirect(url_for('main.doctor_dashboard' if session.get('role')=='doctor' else 'main.dashboard')) if session.get('user_id') else redirect(url_for('main.login'))
@main_bp.route('/login')
def login(): return redirect(url_for('main.dashboard')) if session.get('user_id') else render_template('login.html')
@main_bp.route('/register')
def register(): return redirect(url_for('main.dashboard')) if session.get('user_id') else render_template('register.html')
@main_bp.route('/doctor/login')
def doctor_login():
    # A doctor already signed in should go straight to the doctor portal.
    if session.get('user_id') and session.get('role') == 'doctor':
        return redirect(url_for('main.doctor_dashboard'))
    # If a patient is signed in, clear the patient session before showing the
    # doctor login page so the doctor authentication flow cannot inherit it.
    if session.get('user_id') and session.get('role') != 'doctor':
        session.clear()
    return render_template('doctor_login.html')
@main_bp.route('/doctor/register')
def doctor_register(): return redirect(url_for('main.doctor_dashboard')) if session.get('role')=='doctor' else render_template('doctor_register.html')
@main_bp.route('/logout')
def logout(): session.clear(); return redirect(url_for('main.login'))

@main_bp.route('/dashboard')
@patient_required
def dashboard(): return render_template('dashboard.html')
@main_bp.route('/chat')
@patient_required
def chat(): return render_template('chat.html')
@main_bp.route('/symptoms')
@patient_required
def symptoms(): return render_template('symptoms.html')
@main_bp.route('/diseases')
@patient_required
def diseases(): return render_template('diseases.html')
@main_bp.route('/medications')
@patient_required
def medications(): return render_template('medications.html')
@main_bp.route('/reports')
@patient_required
def reports(): return render_template('reports.html')
@main_bp.route('/timeline')
@patient_required
def timeline(): return render_template('timeline.html')
@main_bp.route('/appointments')
@patient_required
def appointments(): return render_template('appointments.html')
@main_bp.route('/settings')
@patient_required
def settings(): return render_template('settings.html')
@main_bp.route('/about')
@login_required
def about(): return render_template('about.html')
@main_bp.route('/profile')
@login_required
def profile(): return render_template('profile.html')
@main_bp.route('/notifications')
@login_required
def notifications(): return render_template('notifications.html')
@main_bp.route('/feedback')
@login_required
def feedback(): return render_template('feedback.html')
@main_bp.route('/contact')
@login_required
def contact(): return render_template('contact.html')
@main_bp.route('/pill-reminder')
@patient_required
def pill_reminder(): return render_template('pill_reminder.html')

@main_bp.route('/doctor/dashboard')
@doctor_required
def doctor_dashboard(): return render_template('doctor_dashboard.html')
@main_bp.route('/doctor/reports')
@doctor_required
def doctor_reports(): return render_template('doctor_reports.html')
@main_bp.route('/video/<room_id>')
@login_required
def video_call(room_id):
    conn=get_db(); row=conn.execute("SELECT patient_id,doctor_id,status FROM appointments WHERE room_id=?",(room_id,)).fetchone(); conn.close()
    if not row or session['user_id'] not in (row['patient_id'],row['doctor_id']): return error_response('Video room not found or access denied.',403)
    if row['status'] not in ('confirmed','in_progress'):
        return error_response('Video consultation becomes available after the appointment is confirmed.',409)
    return render_template('video_call.html',room_id=room_id)

@api_bp.route('/auth/register',methods=['POST'])
def api_register():
    data=request.get_json(silent=True) or {}; valid,err=validate_json_body(data,['first_name','last_name','email','password'])
    if not valid:return error_response(err)
    if len(str(data['password']))<8:return error_response('Password must be at least 8 characters.')
    try:
        uid=create_user(data['first_name'],data['last_name'],data['email'],data.get('date_of_birth',''),data['password'],'patient')
        return success_response({'registered':True,'user_id':uid},201)
    except Exception as e:
        if 'UNIQUE constraint failed' in str(e):return error_response('An account with this email already exists.',409)
        logger.exception('Registration error');return error_response('Registration failed.',500)

@api_bp.route('/auth/doctor/register',methods=['POST'])
def api_doctor_register():
    data=request.get_json(silent=True) or {}; valid,err=validate_json_body(data,['first_name','last_name','email','password','specialization','license_number'])
    if not valid:return error_response(err)
    if len(str(data['password']))<8:return error_response('Password must be at least 8 characters.')
    try:
        uid=create_user(data['first_name'],data['last_name'],data['email'],'',data['password'],'doctor',data['specialization'],data['license_number'])
        return success_response({'registered':True,'user_id':uid},201)
    except Exception as e:
        if 'UNIQUE constraint failed' in str(e):return error_response('An account with this email already exists.',409)
        logger.exception('Doctor registration error');return error_response('Doctor registration failed.',500)

def _login(data,role):
    valid,err=validate_json_body(data,['email','password'])
    if not valid:return None,error_response(err)
    user=authenticate_user(data['email'],data['password'],role)
    if not user:return None,error_response('Invalid email, password, or account type.',401)
    session.clear();session['user_id']=user['id'];session['user_email']=user['email'];session['role']=user['role'];session['user_name']=f"{user['first_name']} {user['last_name']}".strip()
    record_login(user['id'],user['email'],request.remote_addr,request.headers.get('User-Agent',''))
    return user,None

@api_bp.route('/auth/login',methods=['POST'])
def api_login():
    user,err=_login(request.get_json(silent=True) or {},'patient')
    if err:return err
    return success_response({'authenticated':True,'role':user['role'],'redirect':'/dashboard'})
@api_bp.route('/auth/doctor/login',methods=['POST'])
def api_doctor_login():
    data=request.get_json(silent=True) or {}
    user,err=_login(data,'doctor')
    if err:return err
    # Never let a doctor session fall through to the patient dashboard.
    session['role']='doctor'
    return success_response({
        'authenticated':True,
        'role':'doctor',
        'user':{'id':user['id'],'name':f"{user['first_name']} {user['last_name']}",'specialization':user['specialization']},
        'redirect':url_for('main.doctor_dashboard')
    })

@api_bp.route('/chat',methods=['POST'])
@patient_required
def api_chat():
    data=request.get_json(silent=True) or {}; valid,err=validate_json_body(data,['message'])
    if not valid:return error_response(err)
    message=sanitize_input(data['message']); history=data.get('history',[]); language=data.get('language','en-IN')
    safe_history=[{'role':h.get('role','user'),'content':sanitize_input(h.get('content',''))} for h in history[-10:]]
    add_chat_message(session['user_id'],'user',message,language)
    result=watsonx.chat(message,safe_history,language)
    add_chat_message(session['user_id'],'assistant',result['text'],language)
    return success_response({'message':result['text'],'model':result.get('agent_id',result.get('model_id','')),'tokens':result.get('output_tokens',0),'demo':result.get('demo',False)})

@api_bp.route('/chat/history')
@patient_required
def chat_history(): return success_response({'messages':get_chat_messages(session['user_id'],100)})

@api_bp.route('/chat/export-docx')
@patient_required
def chat_export_docx():
    rows=get_chat_messages(session['user_id'],200)
    if not rows:return error_response('No AI chat history to export.')
    try:
        from docx import Document
        folder=os.path.join(config.UPLOAD_FOLDER,'chat_exports',str(session['user_id']));os.makedirs(folder,exist_ok=True)
        path=os.path.join(folder,f"doctor-talk-chat-{session['user_id']}.docx")
        doc=Document();doc.add_heading('DoctorTalk — AI Chat Record',0);doc.add_paragraph(f"Patient: {session.get('user_name','User')}")
        for r in rows:
            who='Patient' if r['role']=='user' else 'DoctorTalk AI'
            doc.add_heading(f"{who} · {r['created_at']} · {r.get('language','en-IN')}",level=2);doc.add_paragraph(r['content'])
        doc.save(path);return send_file(path,as_attachment=True,download_name=os.path.basename(path),mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    except Exception:
        logger.exception('Chat DOCX export failed');return error_response('Could not create the chat document.',500)

@api_bp.route('/chat/export-pdf')
@patient_required
def chat_export_pdf():
    rows=get_chat_messages(session['user_id'],200)
    if not rows:return error_response('No AI chat history to export.')
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib import colors
        folder=os.path.join(config.UPLOAD_FOLDER,'chat_exports',str(session['user_id']));os.makedirs(folder,exist_ok=True)
        path=os.path.join(folder,f"doctor-talk-chat-{session['user_id']}.pdf")
        styles=getSampleStyleSheet(); title=styles['Title']; title.alignment=TA_CENTER
        doc=SimpleDocTemplate(path,pagesize=A4,rightMargin=42,leftMargin=42,topMargin=42,bottomMargin=42)
        story=[Paragraph('DoctorTalk — AI Chat Record',title),Spacer(1,16),Paragraph(f"Patient: {session.get('user_name','User')}",styles['Normal']),Spacer(1,10)]
        for r in rows:
            who='Patient' if r['role']=='user' else 'DoctorTalk AI'
            txt=str(r['content']).replace('&','&amp;').replace('<','&lt;').replace('>','&gt;').replace('\n','<br/>')
            story += [Paragraph(f"<b>{who}</b> · {r['created_at']} · {r.get('language','en-IN')}",styles['Heading4']),Paragraph(txt,styles['BodyText']),Spacer(1,10)]
        doc.build(story)
        return send_file(path,as_attachment=True,download_name=os.path.basename(path),mimetype='application/pdf')
    except Exception:
        logger.exception('Chat PDF export failed');return error_response('Could not create the chat document.',500)

@api_bp.route('/symptoms',methods=['POST'])
@patient_required
def api_symptoms():
    data=request.get_json(silent=True) or {}; valid,err=validate_json_body(data,['symptoms'])
    if not valid:return error_response(err)
    result=watsonx.check_symptoms(sanitize_input(data['symptoms']));return success_response({'analysis':result['text'],'demo':result.get('demo',False)})
@api_bp.route('/disease',methods=['POST'])
@patient_required
def api_disease():
    data=request.get_json(silent=True) or {}; valid,err=validate_json_body(data,['disease'])
    if not valid:return error_response(err)
    d=sanitize_input(data['disease'],200);result=watsonx.get_disease_info(d);return success_response({'info':result['text'],'disease':d,'demo':result.get('demo',False)})
@api_bp.route('/medication',methods=['POST'])
@patient_required
def api_medication():
    data=request.get_json(silent=True) or {};valid,err=validate_json_body(data,['medication'])
    if not valid:return error_response(err)
    d=sanitize_input(data['medication'],200);result=watsonx.get_medication_info(d);return success_response({'info':result['text'],'medication':d,'demo':result.get('demo',False)})

@api_bp.route('/report/upload',methods=['POST'])
@patient_required
def api_report_upload():
    if 'file' not in request.files:return error_response('No file provided.')
    try:
        filename,filepath=save_upload(request.files['file']);text=extract_text_from_file(filepath)
        if not text.strip():return error_response('Could not extract text from the uploaded file.')
        result=watsonx.analyze_report(text)
        conn=get_db();cur=conn.execute("INSERT INTO medical_reports(patient_id,filename,stored_path,extracted_text,analysis,status) VALUES(?,?,?,?,?,?)",(session['user_id'],filename,filepath,text,result['text'],'ready'));conn.commit();rid=cur.lastrowid;conn.close()
        return success_response({'report_id':rid,'filename':filename,'extracted_text':text[:1000],'analysis':result['text'],'demo':result.get('demo',False)})
    except ValueError as e:return error_response(str(e))
    except Exception:logger.exception('Report upload error');return error_response('Failed to process the file. Please try again.',500)

@api_bp.route('/doctors')
@patient_required
def doctors(): return success_response({'doctors':list_doctors()})
@api_bp.route('/appointments',methods=['GET','POST'])
@patient_required
def patient_appointments():
    conn=get_db()
    if request.method=='POST':
        data=request.get_json(silent=True) or {};valid,err=validate_json_body(data,['doctor_id','title','appointment_time'])
        if not valid:conn.close();return error_response(err)
        try:
            doctor=conn.execute("SELECT id FROM users WHERE id=? AND role='doctor'",(int(data['doctor_id']),)).fetchone()
            if not doctor:
                conn.close();return error_response('Selected doctor does not exist.',400)
            aid,room=create_appointment(session['user_id'],int(data['doctor_id']),sanitize_input(data['title'],200),sanitize_input(data.get('specialty',''),100),sanitize_input(data.get('notes',''),1000),sanitize_input(data['appointment_time'],50))
            conn.close();return success_response({'appointment_id':aid,'room_id':room},201)
        except Exception as e:conn.close();return error_response(str(e),400)
    rows=conn.execute("""SELECT a.*, d.first_name||' '||d.last_name doctor_name,d.specialization doctor_specialty FROM appointments a JOIN users d ON d.id=a.doctor_id WHERE a.patient_id=? ORDER BY a.appointment_time""",(session['user_id'],)).fetchall();conn.close()
    return success_response({'appointments':[dict(r) for r in rows]})

@api_bp.route('/appointments/<int:appointment_id>/cancel',methods=['POST'])
@patient_required
def cancel_appointment(appointment_id):
    conn=get_db();cur=conn.execute("UPDATE appointments SET status='cancelled' WHERE id=? AND patient_id=?",(appointment_id,session['user_id']));conn.commit();conn.close();return success_response({'updated':cur.rowcount>0})

@api_bp.route('/doctor/appointments')
@doctor_required
def doctor_appointments():
    conn=get_db();rows=conn.execute("""SELECT a.*, p.first_name||' '||p.last_name patient_name,p.email patient_email FROM appointments a JOIN users p ON p.id=a.patient_id WHERE a.doctor_id=? ORDER BY CASE a.status WHEN 'requested' THEN 0 WHEN 'confirmed' THEN 1 ELSE 2 END,a.appointment_time""",(session['user_id'],)).fetchall();conn.close();return success_response({'appointments':[dict(r) for r in rows]})
@api_bp.route('/doctor/appointments/<int:appointment_id>',methods=['PATCH'])
@doctor_required
def doctor_update_appointment(appointment_id):
    data=request.get_json(silent=True) or {};status=data.get('status')
    if status not in ('requested','confirmed','completed','cancelled','in_progress'):return error_response('Invalid appointment status.')
    conn=get_db();cur=conn.execute("UPDATE appointments SET status=? WHERE id=? AND doctor_id=?",(status,appointment_id,session['user_id']));conn.commit();conn.close();return success_response({'updated':cur.rowcount>0})

@api_bp.route('/doctor/reports')
@doctor_required
def doctor_report_list():
    conn=get_db();rows=conn.execute("""SELECT r.id,r.filename,r.extracted_text,r.analysis,r.status,r.created_at,p.first_name||' '||p.last_name patient_name,p.email patient_email FROM medical_reports r JOIN users p ON p.id=r.patient_id WHERE EXISTS (SELECT 1 FROM appointments a WHERE a.patient_id=r.patient_id AND a.doctor_id=?) ORDER BY r.created_at DESC""",(session['user_id'],)).fetchall();conn.close();return success_response({'reports':[dict(r) for r in rows]})

# Simple DB-backed WebRTC signaling for a single Flask instance.
@api_bp.route('/video/<room_id>',methods=['GET','POST'])
@login_required
def video_signal(room_id):
    conn=get_db(); appt=conn.execute("SELECT patient_id,doctor_id FROM appointments WHERE room_id=?",(room_id,)).fetchone()
    if not appt or session['user_id'] not in (appt['patient_id'],appt['doctor_id']): conn.close(); return error_response('Video room access denied.',403)
    status=conn.execute('SELECT status FROM appointments WHERE room_id=?',(room_id,)).fetchone()['status']
    if status not in ('confirmed','in_progress'):
        conn.close(); return error_response('Video consultation is not active for this appointment yet.',409)
    conn.execute("INSERT OR IGNORE INTO video_signals(room_id) VALUES(?)",(room_id,))
    if request.method=='POST':
        data=request.get_json(silent=True) or {};kind=data.get('kind');value=data.get('value')
        if kind=='offer':conn.execute('UPDATE video_signals SET offer=? WHERE room_id=?',(json.dumps(value),room_id))
        elif kind=='answer':conn.execute('UPDATE video_signals SET answer=? WHERE room_id=?',(json.dumps(value),room_id))
        elif kind=='caller-candidate':conn.execute("UPDATE video_signals SET caller_candidates=json_insert(caller_candidates,'$[#]',json(?)) WHERE room_id=?",(json.dumps(value),room_id))
        elif kind=='callee-candidate':conn.execute("UPDATE video_signals SET callee_candidates=json_insert(callee_candidates,'$[#]',json(?)) WHERE room_id=?",(json.dumps(value),room_id))
        else:conn.close();return error_response('Unknown signaling message.')
        conn.commit();conn.close();return success_response({'saved':True})
    row=conn.execute('SELECT * FROM video_signals WHERE room_id=?',(room_id,)).fetchone();conn.close()
    d=dict(row) if row else {'offer':None,'answer':None,'caller_candidates':'[]','callee_candidates':'[]'}
    for k in ('offer','answer','caller_candidates','callee_candidates'):
        try:d[k]=json.loads(d[k]) if d[k] else ([] if 'candidate' in k else None)
        except Exception:pass
    return success_response(d)

@api_bp.route('/health-tip',methods=['GET'])
@patient_required
def api_health_tip():result=watsonx.get_health_tip();return success_response({'tip':result['text'],'demo':result.get('demo',False)})
@api_bp.route('/wellness-plan',methods=['POST'])
@patient_required
def api_wellness_plan():
    data=request.get_json(silent=True) or {};profile={k:sanitize_input(str(data.get(k,''))) for k in ('age','gender','conditions','goals')};result=watsonx.generate_wellness_plan(profile);return success_response({'plan':result['text'],'demo':result.get('demo',False)})

@api_bp.route('/session/save-chat',methods=['POST'])
@patient_required
def save_chat():return success_response({'saved':True})
@api_bp.route('/session/load-chat')
@patient_required
def load_chat():return success_response({'messages':get_chat_messages(session['user_id'],100)})
@api_bp.route('/session/clear-chat',methods=['POST'])
@patient_required
def clear_chat():
    conn=get_db();conn.execute('DELETE FROM chat_messages WHERE user_id=?',(session['user_id'],));conn.commit();conn.close();return success_response({'cleared':True})

@main_bp.app_errorhandler(404)
def not_found(e):return render_template('errors/404.html'),404
@main_bp.app_errorhandler(500)
def server_error(e):return render_template('errors/500.html'),500
